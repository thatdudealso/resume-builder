from __future__ import annotations

import io

import pytest
from docx import Document


def _styled_docx_bytes() -> bytes:
    document = Document()
    document.add_paragraph("Original heading", style="Title")
    document.add_paragraph("Original body", style="Intense Quote")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


@pytest.mark.asyncio
async def test_export_all_formats_unlocked(client, monkeypatch):
    await client.post(
        "/api/v1/auth/register",
        json={"email": "formats@test.com", "password": "password123"},
        headers={"X-Device-Fingerprint": "fp"},
    )
    monkeypatch.setattr(
        "apps.web.api.v1.resumes.extract_text_from_upload",
        lambda f, d: "SUMMARY\nEngineer\nEXPERIENCE\nBuilt systems.",
    )
    monkeypatch.setattr("apps.web.api.v1.resumes.upload_bytes", lambda k, d, c: k)
    template_bytes = _styled_docx_bytes()
    upload = await client.post(
        "/api/v1/resumes",
        files={
            "file": (
                "r.docx",
                io.BytesIO(template_bytes),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    resume_id = upload.json()["resume_id"]

    async def bg(run_id, variant=None):
        import packages.db.session as db_session
        from packages.db.models.agent_run import AgentRun

        async with db_session.SessionLocal() as s:
            run = await s.get(AgentRun, run_id)
            if run:
                run.status = "completed"
                run.output_locked = False
                run.is_free_trial_run = False
                run.final_output = {"plain_text": "Full resume text for export."}
                await s.commit()

    monkeypatch.setattr("apps.web.services.run_launcher.execute_run_background", bg)
    run_resp = await client.post(
        "/api/v1/runs",
        json={"resume_id": resume_id, "jd_text": "Python developer " * 5},
    )
    run_id = run_resp.json()["run_id"]
    import asyncio

    await asyncio.sleep(0.05)
    exported: dict[str, bytes] = {}

    def capture_export(key, data, content_type):
        exported[content_type] = data
        return key

    monkeypatch.setattr("apps.web.api.v1.exports.download_bytes", lambda key: template_bytes)
    monkeypatch.setattr("apps.web.api.v1.exports.upload_bytes", capture_export)
    monkeypatch.setattr("apps.web.api.v1.exports.export_pdf", lambda text, **kwargs: b"%PDF-1.4")
    for fmt in ("txt", "docx", "pdf"):
        resp = await client.post("/api/v1/exports", json={"run_id": run_id, "format": fmt})
        assert resp.status_code in (200, 402)

    docx = Document(
        io.BytesIO(
            exported["application/vnd.openxmlformats-officedocument.wordprocessingml.document"]
        )
    )
    assert docx.paragraphs[0].text == "Full resume text for export."
    assert docx.paragraphs[0].style.name == "Title"
