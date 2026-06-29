from __future__ import annotations

import io
import sys

from docx import Document

from packages.export.docx_export import export_docx, export_txt
from packages.export.pdf_export import export_pdf
from packages.export.pdf_ingest import extract_text_from_upload
from packages.integrations.crypto.nowpayments import verify_ipn_signature


def test_export_txt():
    assert export_txt("hello").decode() == "hello"


def test_export_docx():
    data = export_docx("line1\nline2")
    assert data[:2] == b"PK"


def test_export_docx_preserves_template_paragraph_style():
    template = Document()
    template.add_paragraph("Original heading", style="Title")
    template.add_paragraph("Original body", style="Intense Quote")
    source = io.BytesIO()
    template.save(source)

    data = export_docx("Tailored heading\nTailored body", template_bytes=source.getvalue())
    result = Document(io.BytesIO(data))

    assert [p.text for p in result.paragraphs] == ["Tailored heading", "Tailored body"]
    assert result.paragraphs[0].style.name == "Title"
    assert result.paragraphs[1].style.name == "Intense Quote"


def test_export_docx_preserves_table_template():
    template = Document()
    table = template.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.cell(0, 0).text = "Original left"
    table.cell(0, 1).text = "Original right"
    source = io.BytesIO()
    template.save(source)

    data = export_docx("Tailored left\nTailored right", template_bytes=source.getvalue())
    result = Document(io.BytesIO(data))

    assert len(result.tables) == 1
    assert result.tables[0].style.name == "Table Grid"
    assert result.tables[0].cell(0, 0).text == "Tailored left"
    assert result.tables[0].cell(0, 1).text == "Tailored right"


def test_export_pdf_falls_back_without_weasyprint(monkeypatch):
    monkeypatch.setitem(sys.modules, "weasyprint", None)

    data = export_pdf("Tailored resume")

    assert data.startswith(b"%PDF-1.4")


def test_extract_txt():
    text = extract_text_from_upload("resume.txt", b"Hello resume")
    assert "Hello" in text


def test_crypto_signature_local(monkeypatch):
    """When no IPN secret is configured, verify_ipn_signature bypasses HMAC
    and returns True in local/test environments (no secret = bypass path)."""
    monkeypatch.setattr("apps.web.config.settings.nowpayments_ipn_secret", "")
    assert verify_ipn_signature(b"{}", "any") is True
