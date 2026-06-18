from __future__ import annotations

from unittest.mock import MagicMock, patch

import pytest
from docx import Document


@patch("packages.export.pdf_ingest.pdfplumber.open")
def test_extract_pdf(mock_open):
    page = MagicMock()
    page.extract_text.return_value = "PDF resume text"
    mock_open.return_value.__enter__.return_value.pages = [page]
    from packages.export.pdf_ingest import extract_text_from_pdf

    text = extract_text_from_pdf(b"%PDF")
    assert "PDF resume" in text


def test_unsupported_format():
    from packages.export.pdf_ingest import extract_text_from_upload

    with pytest.raises(ValueError):
        extract_text_from_upload("file.doc", b"x")


def test_extract_docx_upload():
    from io import BytesIO

    from packages.export.pdf_ingest import extract_text_from_upload

    document = Document()
    document.add_paragraph("Summary")
    document.add_paragraph("Built reliable APIs.")
    buffer = BytesIO()
    document.save(buffer)

    text = extract_text_from_upload("resume.docx", buffer.getvalue())

    assert text == "Summary\nBuilt reliable APIs."


def test_empty_upload_text_rejected():
    from packages.export.pdf_ingest import extract_text_from_upload

    with pytest.raises(ValueError, match="readable text"):
        extract_text_from_upload("resume.txt", b"   \n\t")
